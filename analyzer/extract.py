"""Turn an uploaded PDF, DOCX or TXT into plain text, plus facts about the file itself."""
from __future__ import annotations

import io
import re

from .errors import AnalyzerError

MAX_BYTES = 5 * 1024 * 1024


def _clean(text: str) -> str:
    text = re.sub(r"\(cid:\d+\)", "•", text)  # symbol-font bullets that PDF readers cannot name
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ").replace("\x0c", "\n")
    text = re.sub(r"[​‌‍﻿]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(data: bytes) -> tuple[str, dict]:
    import pdfplumber

    pages_text, tables, images = [], 0, 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = len(pdf.pages)
            for page in pdf.pages:
                pages_text.append(page.extract_text() or "")
                try:
                    tables += len(page.find_tables())
                except Exception:
                    pass
                images += len(page.images)
    except Exception as exc:
        raise AnalyzerError("This PDF could not be opened. It may be damaged or password protected.") from exc
    return "\n".join(pages_text), {"kind": "pdf", "pages": pages, "tables": tables, "images": images}


def _docx(data: bytes) -> tuple[str, dict]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise AnalyzerError("This Word file could not be opened. Save it as .docx and try again.") from exc

    lines = []
    for section in document.sections:  # contact details often live in the page header
        for p in section.header.paragraphs:
            if p.text.strip():
                lines.append(p.text)
    for p in document.paragraphs:
        style = (p.style.name or "").lower() if p.style is not None else ""
        is_list = style.startswith("list") or p._p.pPr is not None and p._p.pPr.numPr is not None
        lines.append(("• " + p.text) if is_list and p.text.strip() else p.text)
    tables = len(document.tables)
    for table in document.tables:
        for row in table.rows:
            seen = []
            for cell in row.cells:
                if cell.text.strip() and cell.text not in seen:
                    seen.append(cell.text)
            if seen:
                lines.append(" | ".join(seen))
    images = len(document.inline_shapes)
    text = "\n".join(lines)
    # Word does not store pages. Estimate: about 550 words fit on a page.
    pages = max(1, round(len(text.split()) / 550 + 0.25))
    return text, {"kind": "docx", "pages": pages, "tables": tables, "images": images, "pages_estimated": True}


def extract(data: bytes, filename: str) -> tuple[str, dict]:
    if not data:
        raise AnalyzerError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise AnalyzerError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        raw, meta = _pdf(data)
    elif name.endswith(".docx") or data[:2] == b"PK":
        raw, meta = _docx(data)
    elif name.endswith(".doc"):
        raise AnalyzerError("Old .doc files are not supported. Save the resume as .docx or PDF and try again.")
    elif name.endswith((".txt", ".md")) or b"\x00" not in data[:2000]:
        try:
            raw = data.decode("utf-8")
        except UnicodeDecodeError:
            raw = data.decode("latin-1")
        meta = {"kind": "txt", "pages": max(1, round(len(raw.split()) / 550 + 0.25)), "tables": 0,
                "images": 0, "pages_estimated": True}
    else:
        raise AnalyzerError("Upload a PDF, DOCX or TXT resume.")

    text = _clean(raw)
    meta["words"] = len(text.split())
    meta["chars"] = len(text)
    if meta["kind"] == "pdf" and meta["words"] < 30 * max(1, meta["pages"]) and meta["words"] < 120:
        meta["scanned"] = True
    return text, meta
